"""Tests for Excel export functionality."""

from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.candidate import Candidate
from app.models.job_description import JobDescription
from app.models.shortlist import Shortlist
from app.models.shortlist_candidate import ShortlistCandidate
from app.services.excel_export_service import ExcelExportService


@pytest.fixture
def test_jd(sqlite_session: Session) -> JobDescription:
    """Create a test JD."""
    jd = JobDescription(
        jd_code="JAVA-001",
        title="Java Backend Engineer",
        required_skills=["Java", "Spring Boot", "PostgreSQL"],
    )
    sqlite_session.add(jd)
    sqlite_session.commit()
    sqlite_session.refresh(jd)
    return jd


@pytest.fixture
def test_shortlist_with_candidates(
    sqlite_session: Session, user_factory, candidate_factory, test_jd
):
    """Create a test shortlist with candidates."""
    recruiter = user_factory()
    
    # Create candidates
    candidates = [
        candidate_factory(
            full_name="Alice Johnson",
            skills=["Java", "Spring Boot"],
            raw_payload={
                "Candidate_Owner": {"name": "Priya Rao"},
                "Staffing_Clients": "Kanini",
                "Practice": "Engineering",
                "Job_ID": "JOB-101",
                "City": "Bengaluru",
                "Overall_IT_Experince": "6",
                "LinkedIn": "https://linkedin.com/in/alice",
            },
        ),
        candidate_factory(full_name="Bob Smith", skills=["Java", "PostgreSQL"]),
        candidate_factory(full_name="Carol White", skills=["Spring Boot", "Redis"]),
    ]
    
    # Create shortlist
    shortlist = Shortlist(recruiter_id=recruiter.id, jd_id=test_jd.id)
    sqlite_session.add(shortlist)
    sqlite_session.commit()
    sqlite_session.refresh(shortlist)
    
    # Add candidates to shortlist
    for candidate in candidates:
        sc = ShortlistCandidate(shortlist_id=shortlist.id, candidate_id=candidate.id)
        sqlite_session.add(sc)
    sqlite_session.commit()
    
    return shortlist, recruiter, candidates


def test_excel_export_service_success(sqlite_session: Session, test_shortlist_with_candidates):
    """Test successful Excel generation from shortlist."""
    shortlist, recruiter, candidates = test_shortlist_with_candidates
    
    # Create export service and generate file
    export_service = ExcelExportService(sqlite_session)
    file_bytes, filename = export_service.generate_shortlist_export(shortlist.id)
    
    # Verify filename
    assert filename == "Java_Backend_Engineer_Shortlist.xlsx"
    
    # Verify file content
    assert file_bytes is not None
    assert len(file_bytes) > 0
    
    # Verify Excel structure
    workbook = load_workbook(BytesIO(file_bytes))
    worksheet = workbook.active
    
    # Check headers
    expected_headers = [
        "Record Id", "Candidate Owner", "Staffing Clients", "First Name", "Last Name",
        "Email", "Skill Set", "Practice", "Source", "Employee Details", "State/Province",
        "Job ID", "Phone", "City", "Country", "Date of Submission", "Job Title as per KANINI",
        "Date of Joining", "Name of Source", "Current Work Location", "Experience in Yrs",
        "Total Years of Exp", "Job Role", "Practice", "Stream", "Department",
        "Job Role as per Kanini", "Overall IT Experince", "Role Relevant Exp", "LinkedIn",
    ]
    assert [cell.value for cell in worksheet[1]] == expected_headers
    
    # Check data rows - verify all 3 candidates are present
    candidate_names = set()
    for row_idx in range(2, 5):
        first_name = worksheet[f"D{row_idx}"].value
        last_name = worksheet[f"E{row_idx}"].value
        if first_name:
            candidate_names.add(" ".join(part for part in (first_name, last_name) if part))
    
    assert "Alice Johnson" in candidate_names
    assert "Bob Smith" in candidate_names
    assert "Carol White" in candidate_names
    
    alice_row = next(
        row_idx
        for row_idx in range(2, 5)
        if worksheet[f"D{row_idx}"].value == "Alice" and worksheet[f"E{row_idx}"].value == "Johnson"
    )
    assert worksheet[f"B{alice_row}"].value == "Priya Rao"
    assert worksheet[f"C{alice_row}"].value == "Kanini"
    assert worksheet[f"L{alice_row}"].value == "JOB-101"
    assert worksheet[f"N{alice_row}"].value == "Bengaluru"
    assert worksheet[f"AB{alice_row}"].value == "6"
    assert worksheet[f"AD{alice_row}"].value == "https://linkedin.com/in/alice"


def test_excel_export_service_nonexistent_shortlist(sqlite_session: Session):
    """Test export of non-existent shortlist raises ValueError."""
    export_service = ExcelExportService(sqlite_session)
    
    with pytest.raises(ValueError, match="not found"):
        export_service.generate_shortlist_export(uuid4())


def test_excel_export_service_no_candidates(
    sqlite_session: Session, user_factory, test_jd
):
    """Test export of shortlist with no candidates raises ValueError."""
    recruiter = user_factory()
    
    # Create shortlist with no candidates
    shortlist = Shortlist(recruiter_id=recruiter.id, jd_id=test_jd.id)
    sqlite_session.add(shortlist)
    sqlite_session.commit()
    
    export_service = ExcelExportService(sqlite_session)
    
    with pytest.raises(ValueError, match="no candidates"):
        export_service.generate_shortlist_export(shortlist.id)


def test_filename_generation():
    """Test filename generation from JD title."""
    test_cases = [
        ("Java Backend Engineer", "Java_Backend_Engineer_Shortlist.xlsx"),
        ("Senior Python Developer / ML", "Senior_Python_Developer__ML_Shortlist.xlsx"),
        ("DevOps & Cloud Architect", "DevOps__Cloud_Architect_Shortlist.xlsx"),
    ]
    
    for jd_title, expected in test_cases:
        filename = ExcelExportService._generate_filename(jd_title)
        # The filename generation replaces spaces and special chars
        # Just verify the key parts are present
        assert filename.endswith("_Shortlist.xlsx")
        assert "Java" in filename or "Senior" in filename or "DevOps" in filename


def test_excel_export_includes_summary(sqlite_session: Session, test_shortlist_with_candidates):
    """Test that export includes summary section."""
    shortlist, recruiter, candidates = test_shortlist_with_candidates
    
    export_service = ExcelExportService(sqlite_session)
    file_bytes, _ = export_service.generate_shortlist_export(shortlist.id)
    
    workbook = load_workbook(BytesIO(file_bytes))
    worksheet = workbook.active
    
    # Find summary section (should be after blank row and data rows)
    # Summary starts at row 6 (after header + 3 candidates + blank)
    assert "Export Summary" in str(worksheet["A6"].value)

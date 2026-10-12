from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from app.models.candidate_review import CandidateReview
from app.models.proposed_field_change import ProposedFieldChange
from app.models.review_batch import ReviewBatch
from app.models.zoho_field_metadata import ZohoFieldMetadataRecord
from app.services.batch_excel_export_service import BatchExcelExportService


def test_export_approved_changes_combines_batches_and_marks_changes_exported(sqlite_session):
    actor_id = uuid4()
    batch_one = ReviewBatch(batch_number="BATCH-001", batch_size=20, status="COMPLETED")
    batch_two = ReviewBatch(batch_number="BATCH-002", batch_size=20, status="COMPLETED")
    sqlite_session.add_all([batch_one, batch_two])
    sqlite_session.flush()

    review_one = CandidateReview(
        batch_id=batch_one.id,
        candidate_id=uuid4(),
        zoho_candidate_id="zoho-1",
        candidate_name="Asha Sharma",
        approval_status="APPROVED",
    )
    review_two = CandidateReview(
        batch_id=batch_two.id,
        candidate_id=uuid4(),
        zoho_candidate_id="zoho-2",
        candidate_name="Ravi Kumar",
        approval_status="APPROVED",
    )
    sqlite_session.add_all([review_one, review_two])
    sqlite_session.flush()

    change_one = ProposedFieldChange(
        candidate_review_id=review_one.id,
        zoho_field_api_name="Current_Employer",
        zoho_field_display_name="Current Employer",
        proposed_value="Acme",
        change_status="APPROVED",
        sync_status="NOT_SENT",
    )
    change_two = ProposedFieldChange(
        candidate_review_id=review_two.id,
        zoho_field_api_name="Relevant_Experience",
        zoho_field_display_name="Relevant Experience",
        proposed_value=7,
        change_status="APPROVED",
        sync_status="NOT_SENT",
    )
    sqlite_session.add_all([
        change_one,
        change_two,
        ZohoFieldMetadataRecord(
            provider="zoho_recruit",
            module="Candidates",
            api_name="Current_Employer",
            display_label="Current Employer",
            data_type="text",
            raw_metadata={},
        ),
        ZohoFieldMetadataRecord(
            provider="zoho_recruit",
            module="Candidates",
            api_name="Relevant_Experience",
            display_label="Relevant Experience",
            data_type="double",
            raw_metadata={},
        ),
    ])
    sqlite_session.commit()

    file_bytes, filename, exported_count = BatchExcelExportService(sqlite_session).export_approved_changes(actor_id)

    workbook = load_workbook(BytesIO(file_bytes))
    worksheet = workbook["Candidates"]
    assert filename.endswith(".xlsx")
    assert exported_count == 2
    assert worksheet.cell(row=1, column=1).value == "Record Id"
    assert [cell.value for cell in worksheet[1]] == [
        "Record Id",
        "Current_Employer",
        "Relevant_Experience",
    ]
    assert worksheet.cell(row=2, column=1).value == "zoho-1"
    assert worksheet.cell(row=3, column=1).value == "zoho-2"
    assert change_one.sync_status == "EXPORTED"
    assert change_two.sync_status == "EXPORTED"
    assert review_one.write_back_status == "EXPORTED"
    assert review_two.write_back_status == "EXPORTED"
from __future__ import annotations

import html
from io import BytesIO

from app.models.template_spec import TemplateSpec
from app.services.kanini_resume_models import KaniniResumeData


def render_html(resume_data: dict, spec: TemplateSpec) -> str:
    resume = KaniniResumeData.model_validate(resume_data)
    sections = [(name, _section(name, resume, spec)) for name in spec.sections]
    content = [(name, value) for name, value in sections if value]
    return (
        f"<style>{_style(spec)}</style>"
        f'<main class="generated-resume columns-{spec.layout.columns} sidebar-{spec.layout.sidebar_position}">'
        f"{_header(resume, spec)}{_layout(content, spec)}</main>"
    )


def render_pdf(resume_data: dict, spec: TemplateSpec) -> bytes:
    try:
        import pymupdf as fitz
    except ImportError as exc:
        raise RuntimeError("PyMuPDF is required for PDF generation") from exc

    page_rect = fitz.paper_rect("a4" if spec.page.size == "A4" else "letter")
    document_html = f'<!doctype html><html><head><meta charset="utf-8"></head><body>{render_html(resume_data, spec)}</body></html>'
    output = BytesIO()
    story = fitz.Story(document_html)
    writer = fitz.DocumentWriter(output)
    try:
        more = 1
        while more:
            device = writer.begin_page(page_rect)
            more, _ = story.place(page_rect)
            story.draw(device)
            writer.end_page()
    finally:
        writer.close()
    return output.getvalue()


def render_docx(resume_data: dict, spec: TemplateSpec) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    resume = KaniniResumeData.model_validate(resume_data)
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27 if spec.page.size == "A4" else 8.5)
    section.page_height = Inches(11.69 if spec.page.size == "A4" else 11)
    margin = Inches(spec.page.margin_inches)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = margin
    normal = document.styles["Normal"]
    normal.font.name = spec.typography.font_family
    normal.font.size = Pt(spec.typography.base_size_pt)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER if spec.header.layout == "centered" else WD_ALIGN_PARAGRAPH.LEFT
    run = heading.add_run(resume.contact.name or "Candidate Name")
    run.bold = True
    run.font.size = Pt(spec.typography.heading_size_pt)
    run.font.color.rgb = RGBColor.from_string(spec.colors.text.lstrip("#"))
    contacts = [value for value in (resume.contact.email, resume.contact.phone, resume.contact.location) if value]
    if contacts:
        document.add_paragraph(("\n" if spec.header.contact_layout == "stacked" else " | ").join(contacts))

    for name in spec.sections:
        lines = _section_lines(name, resume)
        if not lines:
            continue
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(spec.spacing.section_gap_pt)
        title = paragraph.add_run(name.replace("_", " ").upper())
        title.bold = True
        title.font.size = Pt(spec.typography.heading_size_pt)
        for line in lines:
            document.add_paragraph(line)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _style(spec: TemplateSpec) -> str:
    width = "210mm" if spec.page.size == "A4" else "8.5in"
    height = "297mm" if spec.page.size == "A4" else "11in"
    divider = "none" if spec.spacing.divider_style == "none" else f"1px solid {spec.colors.accent if spec.spacing.divider_style == 'accent' else spec.colors.muted}"
    return (
        "*{box-sizing:border-box}.generated-resume{"
        f"width:min(100%,{width});min-height:{height};margin:0 auto;padding:{spec.page.margin_inches}in;"
        f"background:#fff;color:{spec.colors.text};font-family:{spec.typography.font_family},sans-serif;"
        f"font-size:{spec.typography.base_size_pt}pt;line-height:{spec.spacing.line_height};overflow-wrap:anywhere;}}"
        f".template-header{{text-align:{spec.header.layout};border-bottom:{divider};padding-bottom:{spec.spacing.section_gap_pt}pt;}}"
        f".template-header h1,.template-section h2{{color:{spec.colors.text};font-size:{spec.typography.heading_size_pt}pt;}}"
        f".template-layout,.template-column{{display:grid;gap:{spec.spacing.section_gap_pt}pt;min-width:0;}}"
        ".columns-2 .template-layout{grid-template-columns:minmax(0,.8fr) minmax(0,2fr);grid-template-areas:'sidebar main'}"
        ".sidebar-left .template-sidebar{grid-area:sidebar}.sidebar-left .template-main{grid-area:main}"
        ".sidebar-right .template-layout{grid-template-areas:'main sidebar'}.sidebar-right .template-sidebar{grid-area:sidebar}"
        ".skill-tags{display:flex;flex-wrap:wrap;gap:4pt}.skill-tags span{padding:2pt 5pt;border:1px solid currentColor;border-radius:3pt}"
        "@media(max-width:640px){.columns-2 .template-layout{grid-template-columns:1fr;grid-template-areas:'sidebar' 'main'}}"
    )


def _header(resume: KaniniResumeData, spec: TemplateSpec) -> str:
    contacts = [value for value in (resume.contact.email, resume.contact.phone, resume.contact.location, resume.contact.linkedin, resume.contact.github) if value]
    separator = "<br>" if spec.header.contact_layout == "stacked" else " | "
    return f'<header class="template-header"><h1>{_escape(resume.contact.name or "Candidate Name")}</h1><div>{separator.join(_escape(value) for value in contacts)}</div></header>'


def _layout(content: list[tuple[str, str]], spec: TemplateSpec) -> str:
    if spec.layout.columns == 1:
        return f'<div class="template-layout"><div class="template-column template-main">{"".join(value for _, value in content)}</div></div>'
    sidebar_names = {"skills", "education", "certifications", "achievements"}
    sidebar = "".join(value for name, value in content if name in sidebar_names)
    main = "".join(value for name, value in content if name not in sidebar_names)
    return f'<div class="template-layout"><aside class="template-column template-sidebar">{sidebar}</aside><div class="template-column template-main">{main}</div></div>'


def _section(name: str, resume: KaniniResumeData, spec: TemplateSpec) -> str:
    lines = _section_lines(name, resume)
    if not lines:
        return ""
    if name == "skills" and spec.spacing.skill_style == "tags":
        body = f'<div class="skill-tags">{"".join(f"<span>{_escape(line)}</span>" for line in lines)}</div>'
    else:
        body = "".join(f"<p>{_escape(line)}</p>" for line in lines)
    return f'<section class="template-section"><h2>{_escape(name.replace("_", " ").title())}</h2>{body}</section>'


def _section_lines(name: str, resume: KaniniResumeData) -> list[str]:
    if name == "summary":
        return [line for line in resume.summary.splitlines() if line]
    if name == "skills":
        return [f"{category}: {', '.join(items)}" for category, items in resume.skills.items() if items]
    if name == "experience":
        return [_entry([item.title, item.company_name or item.company, item.dates, item.location], item.responsibilities) for item in resume.experience]
    if name == "projects":
        return [_entry([item.name, item.client, item.role, item.duration], [item.description, *item.responsibilities]) for item in resume.projects]
    if name == "education":
        return [" | ".join(value for value in (item.degree, item.institution, item.year, item.gpa) if value) for item in resume.education]
    if name == "certifications":
        return resume.certifications
    if name == "achievements":
        return resume.achievements
    return []


def _entry(values: list[str], details: list[str]) -> str:
    return "\n".join([" | ".join(value for value in values if value), *(value for value in details if value)]).strip()


def _escape(value: str) -> str:
    return html.escape(str(value or ""))
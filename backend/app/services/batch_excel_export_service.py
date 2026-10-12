from __future__ import annotations

from collections import OrderedDict
from datetime import datetime
from io import BytesIO
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.candidate_review import CandidateReview
from app.models.proposed_field_change import ProposedFieldChange
from app.models.zoho_field_metadata import ZohoFieldMetadataRecord


class BatchExcelExportService:
    def __init__(self, session: Session) -> None:
        self.session = session

    def export_approved_changes(self, actor_id: UUID) -> tuple[bytes, str, int]:
        changes = (
            self.session.query(ProposedFieldChange, CandidateReview)
            .join(CandidateReview, CandidateReview.id == ProposedFieldChange.candidate_review_id)
            .filter(
                ProposedFieldChange.change_status == "APPROVED",
                ProposedFieldChange.sync_status == "NOT_SENT",
            )
            .order_by(ProposedFieldChange.updated_at.desc(), ProposedFieldChange.created_at.desc())
            .all()
        )

        selected: OrderedDict[tuple[str, str], tuple[ProposedFieldChange, CandidateReview]] = OrderedDict()
        for change, review in changes:
            if not review.zoho_candidate_id:
                raise ValueError(f"Candidate review {review.id} does not have a Zoho candidate ID")
            key = (review.zoho_candidate_id, change.zoho_field_api_name)
            selected.setdefault(key, (change, review))

        if not selected:
            raise ValueError("No approved, unexported field changes are available")

        metadata = {
            item.api_name: item
            for item in self.session.query(ZohoFieldMetadataRecord)
            .filter(
                ZohoFieldMetadataRecord.provider == "zoho_recruit",
                ZohoFieldMetadataRecord.module == "Candidates",
            )
            .all()
        }

        rows_by_candidate: OrderedDict[str, dict[str, object]] = OrderedDict()
        selected_changes: list[tuple[ProposedFieldChange, CandidateReview]] = []
        for (zoho_candidate_id, field_api_name), pair in selected.items():
            change, review = pair
            row = rows_by_candidate.setdefault(
                zoho_candidate_id,
                {"Record Id": zoho_candidate_id},
            )
            row[field_api_name] = self._display_value(change.proposed_value)
            selected_changes.append(pair)

        field_names = list(OrderedDict.fromkeys(change.zoho_field_api_name for change, _ in selected_changes))
        headers = ["Record Id", *field_names]

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Candidates"
        worksheet.append(headers)
        self._style_header(worksheet)

        for row in rows_by_candidate.values():
            worksheet.append([row.get(header, "") for header in headers])

        metadata_sheet = workbook.create_sheet("Field Metadata")
        metadata_sheet.append(["API Name", "Display Name", "Data Type"])
        self._style_header(metadata_sheet)
        for field_name in field_names:
            item = metadata.get(field_name)
            metadata_sheet.append([
                field_name,
                item.display_label if item else "",
                item.data_type if item else "",
            ])

        for sheet in (worksheet, metadata_sheet):
            for column_index, column_cells in enumerate(sheet.columns, start=1):
                width = max(len(str(cell.value or "")) for cell in column_cells) + 3
                sheet.column_dimensions[get_column_letter(column_index)].width = min(max(width, 16), 40)

        exported_at = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"zoho_candidate_updates_{exported_at}.xlsx"

        for change, review in selected_changes:
            change.sync_status = "EXPORTED"
            change.sync_error = None
            review.write_back_status = "EXPORTED"
            review.write_back_error = None

        output = BytesIO()
        workbook.save(output)
        return output.getvalue(), filename, len(selected_changes)

    @staticmethod
    def _style_header(sheet) -> None:
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    @staticmethod
    def _display_value(value: object) -> object:
        if isinstance(value, dict):
            return value.get("name") or value.get("Name") or value.get("value") or ""
        if isinstance(value, list):
            return ", ".join(str(item) for item in value if item is not None)
        return value if value is not None else ""